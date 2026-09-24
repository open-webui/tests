"""Dependency smoke: every document format Open WebUI reads, uploaded through the API.

An upload with `process=true` runs through `retrieval/loaders/main.py`, which hands each format
to a third-party library: .pdf to pypdf, .docx to docx2txt, .pptx, .xlsx, .xls, .xml, .rst, .epub
and .odt to unstructured's partitioners (python-pptx; pandas on openpyxl or xlrd behind a
msoffcrypto encryption check; pypandoc and the pandoc binary), .html to BeautifulSoup, plain text
through chardet's encoding hint, and every result through ftfy. With `PDF_EXTRACT_IMAGES` on, a
PDF's images are opened by Pillow and read by rapidocr on onnxruntime and OpenCV. A dependency
bump that breaks one of those paths fails the upload or loses the text, which
`GET /api/v1/files/{id}/data/content` shows. The library contracts are in unit/deps/.

EUC-KR and Shift-JIS text fail on dev until #31352 is fixed (PR #31356): chardet 7.4.3 says
CP949 for EUC-KR, which the codec map in `_detect_text_encoding` lacks, and Shift-JIS is missing
from its try order, so both are decoded as GB18030 mojibake.

Discriminates: passes on dev bbfa876af (.rst, .epub and .odt with a pandoc binary on PATH). One
backend copy broke pypdf's `extract_text`, `docx2txt.process`, the xlsx, rst and epub partitions
and `chardet.detect`; another broke rapidocr's `RapidOCR`, the pptx, xml and odt partitions,
BeautifulSoup's `get_text` and `ftfy.fix_text`. Each copy turned exactly its own formats red and
left the others green. Mapping cp949 in a third copy makes the EUC-KR case pass.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import httpx
import pytest

pytestmark = [pytest.mark.depcheck, pytest.mark.api, pytest.mark.requires_source]

RETRIEVAL_CONFIG = ("/api/v1/retrieval/config", "/api/v1/retrieval/config/update")
FIXTURES = Path(__file__).parent / "fixtures"

MARKER = "harbour lighthouse budget"
SENTENCE = f"The {MARKER} was approved."
OCR_WORD = "LIGHTHOUSE"

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client: httpx.Client, filename: str, content: bytes, content_type: str) -> str:
    """Upload and process a file before answering; returns its id."""
    uploaded = client.post(
        "/api/v1/files/",
        params={"process": "true", "process_in_background": "false"},
        files={"file": (filename, content, content_type)},
    )
    assert uploaded.status_code == 200, uploaded.text
    return uploaded.json()["id"]


def _read_back(client: httpx.Client, file_id: str) -> str:
    stored = client.get(f"/api/v1/files/{file_id}").json()
    assert stored["data"].get("status") == "completed", (
        f"{stored['filename']} was not read: {stored['data'].get('error')}"
    )
    read = client.get(f"/api/v1/files/{file_id}/data/content")
    assert read.status_code == 200, read.text
    return read.json()["content"]


def _upload_and_read(client: httpx.Client, filename: str, content: bytes, content_type: str) -> str:
    return _read_back(client, _upload(client, filename, content, content_type))


# ---------------------------------------------------------------- fixtures built in code


def _saved(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _text_pdf() -> bytes:
    """One page carrying the sentence in Helvetica, the smallest PDF pypdf reads text from."""
    stream = f"BT /F1 12 Tf 72 720 Td ({SENTENCE}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref_offset = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1)
    pdf += b"startxref\n%d\n%%%%EOF\n" % xref_offset
    return bytes(pdf)


def _image_pdf(word: str) -> bytes:
    """A scan: the word drawn into an image, saved as a PDF with no text layer."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (900, 200), "white")
    font = ImageFont.load_default(size=64)
    ImageDraw.Draw(image).text((30, 60), word, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PDF")
    return buffer.getvalue()


def _docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(SENTENCE)
    return _saved(document)


def _pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Harbour review"
    slide.placeholders[1].text = SENTENCE
    return _saved(presentation)


def _xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.append(["item", "note"])
    workbook.active.append(["lighthouse", MARKER])
    return _saved(workbook)


def _xls() -> bytes:
    # xlwt, the only .xls writer, is no dependency; LibreOffice saved the rows `_xlsx` writes.
    return (FIXTURES / "budget.xls").read_bytes()


def _xml() -> bytes:
    return f"<?xml version='1.0'?><notes><note>{SENTENCE}</note></notes>".encode()


def _html() -> bytes:
    page = f"<html><head><title>Harbour</title></head><body><p>{SENTENCE}</p></body></html>"
    return page.encode()


def _rst() -> bytes:
    return f"Harbour\n=======\n\n{SENTENCE}\n".encode()


def _container(mimetype: str, entries: dict[str, str]) -> bytes:
    """An EPUB or OpenDocument zip: the `mimetype` entry first and uncompressed."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", mimetype, compress_type=zipfile.ZIP_STORED)
        for name, text in entries.items():
            archive.writestr(name, text)
    return buffer.getvalue()


def _epub() -> bytes:
    container = (
        '<?xml version="1.0"?>'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/>'
        "</rootfiles></container>"
    )
    package = (
        '<?xml version="1.0"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="id">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Harbour</dc:title>'
        '<dc:identifier id="id">harbour</dc:identifier><dc:language>en</dc:language></metadata>'
        '<manifest><item id="text" href="text.xhtml" media-type="application/xhtml+xml"/>'
        '</manifest><spine><itemref idref="text"/></spine></package>'
    )
    chapter = (
        '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Harbour</title></head>'
        f"<body><p>{SENTENCE}</p></body></html>"
    )
    return _container(
        "application/epub+zip",
        {"META-INF/container.xml": container, "content.opf": package, "text.xhtml": chapter},
    )


def _odt() -> bytes:
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"'
        ' manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/"'
        ' manifest:media-type="application/vnd.oasis.opendocument.text"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' office:version="1.2"><office:styles/></office:document-styles>'
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        f"<office:body><office:text><text:p>{SENTENCE}</text:p></office:text></office:body>"
        "</office:document-content>"
    )
    return _container(
        "application/vnd.oasis.opendocument.text",
        {"META-INF/manifest.xml": manifest, "content.xml": content, "styles.xml": styles},
    )


def _require_pandoc() -> None:
    import pypandoc

    try:
        pypandoc.get_pandoc_version()
    except OSError:
        pytest.skip("no pandoc binary; unstructured converts .rst, .epub and .odt with it")


# ---------------------------------------------------------------- one upload per format


@pytest.mark.parametrize(
    ("filename", "build", "content_type"),
    [
        pytest.param("notes.pdf", _text_pdf, "application/pdf", id="pdf"),
        pytest.param("notes.docx", _docx, DOCX_TYPE, id="docx"),
        pytest.param("slides.pptx", _pptx, PPTX_TYPE, id="pptx"),
        pytest.param("budget.xlsx", _xlsx, XLSX_TYPE, id="xlsx"),
        pytest.param("budget.xls", _xls, "application/vnd.ms-excel", id="xls"),
        pytest.param("notes.xml", _xml, "application/xml", id="xml"),
        pytest.param("notes.html", _html, "text/html", id="html"),
    ],
)
def test_an_uploaded_document_is_read_as_text(make_user, filename, build, content_type):
    with make_user().client() as client:
        content = _upload_and_read(client, filename, build(), content_type)

    assert MARKER in content, f"{filename} was read as {content!r}"
    assert "<" not in content, f"{filename} kept its markup: {content!r}"


@pytest.mark.parametrize(
    ("filename", "build", "content_type"),
    [
        pytest.param("notes.rst", _rst, "text/x-rst", id="rst"),
        pytest.param("notes.epub", _epub, "application/epub+zip", id="epub"),
        pytest.param("notes.odt", _odt, "application/vnd.oasis.opendocument.text", id="odt"),
    ],
)
def test_a_document_pandoc_converts_is_read_as_text(make_user, filename, build, content_type):
    _require_pandoc()
    with make_user().client() as client:
        content = _upload_and_read(client, filename, build(), content_type)

    assert MARKER in content, f"{filename} was read as {content!r}"


# ---------------------------------------------------------------- text encodings


@pytest.mark.parametrize(
    ("codec", "text"),
    [
        pytest.param("gb18030", "港口灯塔的预算已经批准。简体中文编码测试。", id="gb18030"),
        pytest.param("big5", "港口燈塔的預算已經批准。繁體中文編碼測試。", id="big5"),
        pytest.param(
            "euc-kr", "항구 등대 예산이 승인되었습니다. 한국어 인코딩 감지 테스트.", id="euc-kr"
        ),
        pytest.param(
            "shift_jis",
            "港の灯台の予算が承認されました。日本語エンコーディング検出テスト。",
            id="shift-jis",
        ),
    ],
)
def test_a_cjk_text_file_is_decoded(make_user, codec, text):
    document = text * 8
    with make_user().client() as client:
        content = _upload_and_read(client, "notes.txt", document.encode(codec), "text/plain")

    assert content == document, f"{codec} was decoded as {content[:40]!r}"


def test_mojibake_is_repaired_and_a_literal_entity_is_kept(make_user):
    # UTF-8 text once decoded as Windows-1252, so its apostrophe became "â€™".
    garbled = "The harbour lighthouse budget wasnâ€™t cut &amp; the keeper stays."
    with make_user().client() as client:
        content = _upload_and_read(client, "notes.txt", garbled.encode(), "text/plain")

    assert "budget wasn't cut" in content, f"the mojibake survived: {content!r}"
    assert "&amp;" in content, f"the literal entity was unescaped: {content!r}"


# ---------------------------------------------------------------- OCR of PDF images


@pytest.fixture
def retrieval_settings(preserve, admin):
    """`update(**settings)` changes the document settings for this test."""
    preserve(RETRIEVAL_CONFIG)
    client = admin.client()

    def update(**settings) -> None:
        updated = client.post(RETRIEVAL_CONFIG[1], json=settings)
        assert updated.status_code == 200, updated.text

    yield update
    client.close()


def test_the_text_in_a_pdf_image_is_read(retrieval_settings, make_user):
    retrieval_settings(PDF_EXTRACT_IMAGES=True)
    with make_user().client() as client:
        content = _upload_and_read(client, "scan.pdf", _image_pdf(OCR_WORD), "application/pdf")

    assert OCR_WORD in content, f"OCR read {content!r}"


def test_without_image_extraction_a_scan_yields_no_text(retrieval_settings, make_user):
    retrieval_settings(PDF_EXTRACT_IMAGES=False)
    with make_user().client() as client:
        file_id = _upload(client, "scan.pdf", _image_pdf(OCR_WORD), "application/pdf")
        read = client.get(f"/api/v1/files/{file_id}/data/content")

    assert OCR_WORD not in read.json()["content"]
